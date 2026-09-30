"""Phase 10 Excel export and presentation for Phase 9 report datasets."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HIGH_PRECISION_DIGITS = 15
UNKNOWN_CURRENCY_BUCKET = "UNKNOWN / UNSPECIFIED CURRENCY"

CLIENT_DETAIL_COLUMN_MAP = {
    "Reporting Month": "Reporting Month",
    "Client": "Client",
    "DevOps Market": "DevOps Market",
    "Normalized DevOps Market": "Normalized DevOps Market",
    "Matched Finance Market Code": "Matched Finance Market",
    "Customer WD": "Customer WD",
    "Match Status": "Match Status",
    "Match Confidence": "Confidence",
    "Match Method": "Match Method",
    "SLA Breach %": "SLA Breach %",
    "Applied Rate": "Applied Rate",
    "Base Credit": "Base Credit",
    "Currency": "Currency",
    "Calculation Rule": "Rule",
    "Calculated Credit": "Calculated Credit",
    "Calculation Status": "Calculation Status",
    "Exception Codes": "Exception",
    "Exception Code": "Primary Exception Code",
    "Explanation / Audit Note": "Explanation",
    "Candidate Count": "Candidate Count",
    "Candidate Customer WD(s)": "Candidate Customer WD(s)",
    "Candidate Market Code(s)": "Candidate Market Code(s)",
    "Candidate Base Credit(s)": "Candidate Base Credit(s)",
    "Candidate Currency/Currencies": "Candidate Currency/Currencies",
    "DevOps Source File": "DevOps Source File",
    "DevOps Source Sheet": "DevOps Source Sheet",
    "DevOps Source Row": "DevOps Source Row",
    "Finance Source File": "Finance Source File",
    "Finance Source Sheet": "Finance Source Sheet",
    "Finance Source Row": "Finance Source Row",
    "Processing Run ID": "Processing Run ID",
    "Source Lineage": "Source Lineage",
    "Currency Limitation": "Currency Limitation",
}

METHODOLOGY_ROWS = [
    (
        "Standard formula",
        "Calculated Credit = Base Credit x MIN(raw SLA Breach %, 20%). SLA percentages are stored as decimal fractions.",
    ),
    (
        "CLIENT-004 special rule",
        "Calculated Credit = 1000 / raw SLA breach. The standard cap and Finance Base Credit do not apply. Zero breach produces zero without division.",
    ),
    (
        "Match confidence",
        "HIGH denotes a unique authoritative Phase 6 match. LOW denotes ambiguous, unmatched, or otherwise non-authoritative matching; confidence does not indicate calculation quality.",
    ),
    (
        "Ambiguous and unmatched values",
        "No candidate is selected or guessed. Standard credit remains blank when a suitable unique HIGH Finance match is unavailable.",
    ),
    (
        "CLIENT-004 currency limitation",
        "Special-rule credits with no unique Finance match have no authoritative currency. They are grouped only in UNKNOWN / UNSPECIFIED CURRENCY, never combined with USD or other currencies.",
    ),
    (
        "MAY26 month conflict",
        "The 66 CLIENT-002 MAY26 rows remain unresolved and ineligible. They are shown in Exceptions / Pending Month Conflict and excluded from authoritative monthly totals.",
    ),
    (
        "Null and zero values",
        "Blank financial cells are intentional null/not-applicable values. A valid zero credit remains numeric zero and is marked ZERO_CREDIT.",
    ),
    (
        "Decimal precision",
        "Phase 9 Decimal amounts are not rounded. Values with more than 15 significant digits are stored as text because Excel numeric cells cannot reliably preserve them. Number formats affect display only.",
    ),
    (
        "Report scope",
        "These workbooks present Phase 9 datasets. They do not recalculate credits, rematch markets, resolve exceptions, convert currencies, or invoke AI.",
    ),
]


@dataclass(frozen=True)
class ExcelExportBundle:
    """Generated workbook paths and the exact Phase 9 frames exported per sheet."""

    paths: dict[str, Path]
    expected_sheets: dict[str, dict[str, pd.DataFrame]]


@dataclass(frozen=True)
class ExcelVerification:
    """Reopen-and-compare checks for generated Excel deliverables."""

    checks: dict[str, bool]
    sheet_names: dict[str, tuple[str, ...]]
    row_counts: dict[str, dict[str, int]]

    @property
    def passed(self) -> bool:
        return all(self.checks.values())


def _is_missing(value: Any) -> bool:
    if value is None:
        return True
    try:
        missing = pd.isna(value)
        return bool(missing) if not hasattr(missing, "__len__") else False
    except (TypeError, ValueError):
        return False


def _significant_digit_count(value: Decimal) -> int:
    digits = value.as_tuple().digits
    significant = len(digits)
    while significant > 1 and digits[significant - 1] == 0:
        significant -= 1
    return significant


def _excel_value(value: Any) -> Any:
    """Keep long Decimal values exact as text; leave null cells blank."""
    if _is_missing(value):
        return None
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        if _significant_digit_count(value) > HIGH_PRECISION_DIGITS:
            return format(value, "f")
        return value
    if isinstance(value, (tuple, list)):
        return "; ".join(str(item) for item in value)
    if hasattr(value, "item"):
        try:
            return _excel_value(value.item())
        except (ValueError, AttributeError):
            return str(value)
    return value


def _excel_frame(frame: pd.DataFrame) -> pd.DataFrame:
    exported = frame.copy(deep=True)
    for column in exported.columns:
        exported[column] = exported[column].map(_excel_value)
    return exported


def _methodology_frame() -> pd.DataFrame:
    return pd.DataFrame(METHODOLOGY_ROWS, columns=["Topic", "Methodology"])


def _select_client_detail(detailed: pd.DataFrame, client: str, month: str) -> pd.DataFrame:
    required = set(CLIENT_DETAIL_COLUMN_MAP)
    missing = sorted(required.difference(detailed.columns))
    if missing:
        raise ValueError(f"Detailed report is missing Phase 9 columns: {missing}")
    filtered = detailed.loc[
        detailed["Client"].eq(client)
        & detailed["Reporting Month"].eq(month)
        & detailed["Reporting Month Status"].eq("RESOLVED")
    ]
    selected = filtered.loc[:, list(CLIENT_DETAIL_COLUMN_MAP)].copy()
    selected.columns = list(CLIENT_DETAIL_COLUMN_MAP.values())
    return selected.reset_index(drop=True)


def _write_workbook(
    path: Path,
    sheets: dict[str, pd.DataFrame],
    titles: dict[str, str],
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            _excel_frame(frame).to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                startrow=1,
            )

    workbook = load_workbook(path)
    workbook.properties.title = path.stem.replace("_", " ")
    workbook.properties.subject = "Credit report Phase 10 presentation export"
    workbook.properties.creator = "Credit Report Agent"
    for sheet_name, frame in sheets.items():
        worksheet = workbook[sheet_name]
        _restore_decimal_cells(worksheet, frame)
        _format_worksheet(
            worksheet,
            frame,
            titles.get(sheet_name, sheet_name),
        )
    workbook.save(path)


def _restore_decimal_cells(worksheet: Any, frame: pd.DataFrame) -> None:
    for row_index, values in enumerate(frame.itertuples(index=False, name=None), start=3):
        for column_index, value in enumerate(values, start=1):
            if isinstance(value, Decimal):
                worksheet.cell(row_index, column_index).value = _excel_value(value)


def _format_worksheet(worksheet: Any, frame: pd.DataFrame, title: str) -> None:
    dark_fill = PatternFill("solid", fgColor="234E52")
    title_fill = PatternFill("solid", fgColor="173F5F")
    status_fills = {
        "CALCULATED": PatternFill("solid", fgColor="E8F3EC"),
        "ZERO_CREDIT": PatternFill("solid", fgColor="EAF2F8"),
        "EXCEPTION": PatternFill("solid", fgColor="FFF2CC"),
        "INELIGIBLE": PatternFill("solid", fgColor="FCE4D6"),
    }
    exception_fills = {
        "MONTH_CONFLICT": PatternFill("solid", fgColor="FCE4D6"),
        "AMBIGUOUS_MATCH": PatternFill("solid", fgColor="FFF2CC"),
        "NO_FINANCE_MATCH": PatternFill("solid", fgColor="FFF2CC"),
        "CURRENCY_UNKNOWN": PatternFill("solid", fgColor="DDEBF7"),
    }
    last_column = max(1, len(frame.columns))
    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )
    title_cell = worksheet.cell(1, 1)
    title_cell.value = title
    title_cell.font = Font(name="Aptos Display", size=15, bold=True, color="FFFFFF")
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 29
    worksheet.row_dimensions[2].height = 34
    worksheet.freeze_panes = "A3"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 90
    worksheet.print_title_rows = "1:2"

    for cell in worksheet[2][:last_column]:
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = dark_fill
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)

    headers = {str(cell.value): cell.column for cell in worksheet[2][:last_column]}
    wrapped_columns = {
        column
        for header, column in headers.items()
        if any(
            token in header.lower()
            for token in ("explanation", "lineage", "candidate", "methodology", "limitation", "exception")
        )
    }
    financial_headers = {
        "Base Credit",
        "Calculated Credit",
        "Total Calculated Credit",
    }
    percent_headers = {"SLA Breach %", "Applied Rate"}
    status_column = headers.get("Calculation Status") or headers.get("Status")
    exception_columns = {
        headers[header]
        for header in ("Exception Code", "Exception Codes", "Exception", "Primary Exception Code")
        if header in headers
    }
    first_data_row = 3
    last_row = max(worksheet.max_row, 2)
    for row_index in range(first_data_row, last_row + 1):
        worksheet.row_dimensions[row_index].height = 30
        status = (
            str(worksheet.cell(row_index, status_column).value)
            if status_column is not None
            else ""
        )
        row_fill = status_fills.get(status)
        for column_index in range(1, last_column + 1):
            cell = worksheet.cell(row_index, column_index)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column_index in wrapped_columns,
            )
            if row_fill is not None:
                cell.fill = row_fill
            if cell.value is None and str(worksheet.cell(2, column_index).value) in (
                financial_headers | percent_headers | {"Currency"}
            ):
                cell.font = Font(name="Aptos", size=10, italic=True, color="7F7F7F")
            header = str(worksheet.cell(2, column_index).value)
            if header in percent_headers and cell.data_type == "n":
                cell.number_format = "0.00%"
            elif header in financial_headers and cell.data_type == "n":
                cell.number_format = '#,##0.################;[Red](#,##0.################);0'
            if header in financial_headers and isinstance(cell.value, str):
                if cell.value and _looks_like_decimal(cell.value):
                    cell.comment = Comment(
                        "Exact high-precision Decimal value stored as text to avoid Excel's 15-digit numeric precision limit.",
                        "Credit Report Agent",
                    )
            if column_index in exception_columns:
                code_text = str(cell.value or "")
                for code, fill in exception_fills.items():
                    if code in code_text:
                        cell.fill = fill
                        break

    for column_index, header in enumerate(frame.columns, start=1):
        values = [str(header)]
        values.extend(
            str(value)
            for value in frame[header].tolist()
            if not _is_missing(value)
        )
        max_length = max((len(value) for value in values), default=12)
        column_width = min(max(max_length + 2, 12), 52)
        if any(token in str(header).lower() for token in ("explanation", "lineage", "methodology")):
            column_width = 48 if str(header) != "Methodology" else 88
        elif "candidate" in str(header).lower() or "exception" in str(header).lower():
            column_width = min(max(column_width, 24), 46)
        worksheet.column_dimensions[get_column_letter(column_index)].width = column_width

    if last_column > 0:
        worksheet.auto_filter.ref = (
            f"A2:{get_column_letter(last_column)}{max(2, worksheet.max_row)}"
        )
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.tabColor = "5B8C85"


def _looks_like_decimal(value: str) -> bool:
    try:
        return Decimal(value).is_finite()
    except Exception:
        return False


def _decimal_value(value: Any) -> Decimal | None:
    if _is_missing(value):
        return None
    try:
        number = Decimal(str(value))
    except Exception:
        return None
    return number if number.is_finite() else None


def _is_positive(value: Any) -> bool:
    number = _decimal_value(value)
    return number is not None and number > 0


def _summary_currency_and_total(
    client_detail: pd.DataFrame,
    client_totals: pd.DataFrame,
    client: str,
    reporting_month: str,
) -> tuple[str, Decimal | None, str]:
    client_groups = client_totals.loc[
        client_totals["Client"].eq(client)
        & client_totals["Reporting Month"].eq(reporting_month)
    ]
    positive_buckets = set(
        client_detail.loc[
            client_detail["Calculated Credit"].map(_is_positive),
            "Reporting Currency Bucket",
        ].dropna()
    )
    if len(positive_buckets) == 1:
        currency_bucket = next(iter(positive_buckets))
    elif len(positive_buckets) > 1:
        return (
            "MULTIPLE CURRENCIES - SEE EXECUTIVE REPORT",
            None,
            "Multiple positive-credit currencies exist; no mixed-currency client total is shown.",
        )
    elif not client_groups.empty:
        calculated_groups = client_groups.loc[client_groups["Calculated Record Count"].gt(0)]
        currency_bucket = (
            calculated_groups.iloc[0]["Currency"]
            if not calculated_groups.empty
            else client_groups.iloc[0]["Currency"]
        )
    else:
        return "UNKNOWN / UNSPECIFIED", None, "No calculated amount is available."

    matches = client_groups.loc[client_groups["Currency"].eq(currency_bucket)]
    if matches.empty:
        return str(currency_bucket), None, "No total exists for this currency bucket."
    total = _decimal_value(matches.iloc[0]["Total Calculated Credit"])
    if currency_bucket == UNKNOWN_CURRENCY_BUCKET:
        status = "UNKNOWN / UNSPECIFIED"
    else:
        status = str(currency_bucket)
    return status, total, (
        "Special-rule credits have no authoritative currency."
        if currency_bucket == UNKNOWN_CURRENCY_BUCKET
        else ""
    )


def _client_summary_data(
    detailed: pd.DataFrame,
    client_totals: pd.DataFrame,
    client: str,
    reporting_month: str,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    client_rows = detailed.loc[
        detailed["Client"].eq(client)
        & detailed["Reporting Month"].eq(reporting_month)
        & detailed["Reporting Month Status"].eq("RESOLVED")
    ]
    client_detail = _select_client_detail(detailed, client, reporting_month)
    breached = client_detail.loc[client_detail["SLA Breach %"].map(_is_positive)].copy()
    breached_records = client_rows.loc[
        client_rows["SLA Breach %"].map(_is_positive)
    ].copy()
    currency_status, exact_total, currency_note = _summary_currency_and_total(
        client_rows, client_totals, client, reporting_month
    )
    rules = tuple(
        dict.fromkeys(
            str(rule)
            for rule in client_rows["Calculation Rule"].tolist()
            if not _is_missing(rule)
        )
    )
    successful_breached = breached.loc[
        breached["Calculation Status"].isin({"CALCULATED", "ZERO_CREDIT"})
        & breached["Calculated Credit"].notna()
    ]

    if exact_total is None:
        display_total = "See currency-specific totals"
    else:
        rounded_total = exact_total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        display_total = f"{rounded_total:,.2f}"
        if currency_status == "UNKNOWN / UNSPECIFIED":
            display_total += " - Currency Unspecified"
        elif currency_status not in {
            "MULTIPLE CURRENCIES - SEE EXECUTIVE REPORT",
            "UNKNOWN / UNSPECIFIED",
        }:
            display_total = f"{currency_status} {display_total}"

    metrics = [
        ("Client", client),
        ("Reporting Month", reporting_month),
        ("Currency / Currency Status", currency_status),
        ("Total Confirmed Calculated Credit", display_total),
        ("Authoritative Total (Exact)", exact_total),
        ("Total Markets Reviewed", len(client_detail)),
        ("Markets with SLA Breach", len(breached)),
        ("Successfully Calculated Breached Markets", len(successful_breached)),
        (
            "Ambiguous Breached Markets",
            int(breached["Match Status"].eq("AMBIGUOUS").sum()),
        ),
        (
            "Unmatched Breached Markets",
            int(breached["Match Status"].eq("UNMATCHED").sum()),
        ),
        ("Applicable Calculation Rule", "; ".join(rules)),
        (
            "Confidence vs Calculation",
            "Match confidence rates Finance-link reliability only. Calculation status is separate; CLIENT-004 may be LOW / UNMATCHED and still CALCULATED by its special rule.",
        ),
    ]
    if currency_note:
        metrics.append(("Currency Limitation", currency_note))
    summary = pd.DataFrame(metrics, columns=["Summary Metric", "Summary Value"])

    breach_columns = [
        "DevOps Market",
        "Matched Finance Market Code",
        "Customer WD",
        "SLA Breach %",
        "Applied Rate",
        "Base Credit",
        "Currency",
        "Calculated Credit",
        "Match Confidence",
        "Match Status",
        "Calculation Status",
        "Exception Codes",
        "Explanation / Audit Note",
    ]
    breached_table = breached_records.loc[:, breach_columns].copy()
    breached_table = breached_table.rename(
        columns={
            "Matched Finance Market Code": "Matched Finance Market",
            "Calculated Credit": "Calculated Credit Receivable",
            "Exception Codes": "Exception",
            "Explanation / Audit Note": "Explanation",
        }
    )
    return summary, breached_table


def _write_client_summary_sheet(
    worksheet: Any,
    summary: pd.DataFrame,
    breached: pd.DataFrame,
    client: str,
) -> None:
    title_fill = PatternFill("solid", fgColor="173F5F")
    section_fill = PatternFill("solid", fgColor="234E52")
    header_fill = PatternFill("solid", fgColor="477C78")
    status_fills = {
        "CALCULATED": PatternFill("solid", fgColor="E8F3EC"),
        "ZERO_CREDIT": PatternFill("solid", fgColor="EAF2F8"),
        "EXCEPTION": PatternFill("solid", fgColor="FFF2CC"),
        "INELIGIBLE": PatternFill("solid", fgColor="FCE4D6"),
    }
    exception_fill = PatternFill("solid", fgColor="FFF2CC")

    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 90
    worksheet.merge_cells("A1:B1")
    title = worksheet["A1"]
    title.value = f"{client} - APRIL 2026 EXECUTIVE SUMMARY"
    title.font = Font(name="Aptos Display", size=15, bold=True, color="FFFFFF")
    title.fill = title_fill
    title.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells("A2:B2")
    worksheet["A2"] = "CLIENT SUMMARY"
    worksheet["A2"].font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
    worksheet["A2"].fill = section_fill
    worksheet["A2"].alignment = Alignment(vertical="center")
    for column, heading in enumerate(summary.columns, start=1):
        cell = worksheet.cell(3, column, heading)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_index, values in enumerate(summary.itertuples(index=False, name=None), start=4):
        worksheet.row_dimensions[row_index].height = 31
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row_index, column_index, _excel_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=column_index == 2)
            if values[0] == "Total Confirmed Calculated Credit":
                cell.font = Font(name="Aptos", bold=True, size=11, color="173F5F")
            if values[0] == "Authoritative Total (Exact)" and isinstance(cell.value, str):
                cell.comment = Comment(
                    "Exact Phase 9 Decimal total retained as text where Excel numeric precision is insufficient.",
                    "Credit Report Agent",
                )
            if values[0] == "Currency Limitation":
                cell.fill = PatternFill("solid", fgColor="DDEBF7")

    section_row = 4 + len(summary) + 1
    worksheet.merge_cells(
        start_row=section_row,
        start_column=1,
        end_row=section_row,
        end_column=len(breached.columns),
    )
    section_cell = worksheet.cell(section_row, 1, "BREACHED MARKETS ONLY (SLA BREACH % > 0)")
    section_cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
    section_cell.fill = section_fill
    section_cell.alignment = Alignment(vertical="center")
    header_row = section_row + 1
    for column_index, heading in enumerate(breached.columns, start=1):
        cell = worksheet.cell(header_row, column_index, heading)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[header_row].height = 34

    headers = {name: index for index, name in enumerate(breached.columns, start=1)}
    for row_offset, values in enumerate(breached.itertuples(index=False, name=None), start=1):
        row_index = header_row + row_offset
        worksheet.row_dimensions[row_index].height = 42
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row_index, column_index, _excel_value(value))
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column_index == headers["Explanation"],
            )
            if column_index in {headers["SLA Breach %"], headers["Applied Rate"]}:
                cell.number_format = "0.00%"
            if column_index in {
                headers["Base Credit"],
                headers["Calculated Credit Receivable"],
            } and cell.data_type == "n":
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
            if headers["Calculation Status"] == column_index:
                status_fill = status_fills.get(str(cell.value))
                if status_fill:
                    cell.fill = status_fill
            if column_index == headers["Match Confidence"]:
                confidence_fills = {
                    "HIGH": PatternFill("solid", fgColor="E8F3EC"),
                    "MEDIUM": PatternFill("solid", fgColor="FFF2CC"),
                    "LOW": PatternFill("solid", fgColor="FCE4D6"),
                }
                if cell.value in confidence_fills:
                    cell.fill = confidence_fills[cell.value]
            if column_index == headers["Exception"] and cell.value:
                cell.fill = exception_fill

    last_row = header_row + len(breached)
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(breached.columns))}{max(header_row, last_row)}"
    )
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.print_title_rows = f"1:{header_row}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.tabColor = "5B8C85"

    widths = {
        "A": 38,
        "B": 24,
        "C": 22,
        "D": 15,
        "E": 14,
        "F": 18,
        "G": 14,
        "H": 24,
        "I": 17,
        "J": 16,
        "K": 19,
        "L": 34,
        "M": 60,
    }
    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 72
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def add_client_executive_summaries(
    detailed: pd.DataFrame,
    client_totals: pd.DataFrame,
    output_dir: str | Path,
    *,
    reporting_month: str = "2026-04",
) -> dict[str, Path]:
    """Insert/update the client-specific summary sheet without rebuilding other reports."""
    directory = Path(output_dir)
    paths: dict[str, Path] = {}
    for client in ("CLIENT-001", "CLIENT-002", "CLIENT-004"):
        filename = f"{client}_APR26_Credit_Report.xlsx"
        path = directory / filename
        if not path.exists():
            raise FileNotFoundError(f"Existing client workbook is required: {path}")
        summary, breached = _client_summary_data(
            detailed, client_totals, client, reporting_month
        )
        workbook = load_workbook(path)
        if "Executive Summary" in workbook.sheetnames:
            del workbook["Executive Summary"]
        worksheet = workbook.create_sheet("Executive Summary", index=0)
        _write_client_summary_sheet(worksheet, summary, breached, client)
        workbook.save(path)
        workbook.close()
        paths[filename] = path
    return paths


def verify_client_executive_summaries(
    paths: dict[str, Path],
    detailed: pd.DataFrame,
    *,
    reporting_month: str = "2026-04",
) -> dict[str, dict[str, Any]]:
    """Verify first-sheet order, breached-row counts/confidence, and Phase 9 values."""
    expected_breached_counts = {
        "CLIENT-001": 7,
        "CLIENT-002": 8,
        "CLIENT-004": 4,
    }
    checks: dict[str, dict[str, Any]] = {}
    for client, breached_count in expected_breached_counts.items():
        filename = f"{client}_APR26_Credit_Report.xlsx"
        path = paths[filename]
        workbook = load_workbook(path, data_only=False)
        worksheet = workbook["Executive Summary"]
        detail_sheet = workbook["Credit Detail"]
        expected_detail = detailed.loc[
            detailed["Client"].eq(client)
            & detailed["Reporting Month"].eq(reporting_month)
            & detailed["Reporting Month Status"].eq("RESOLVED")
        ]
        expected_detail_export = _select_client_detail(detailed, client, reporting_month)
        summary_rows = {
            worksheet.cell(row, 1).value: worksheet.cell(row, 2).value
            for row in range(4, worksheet.max_row + 1)
            if worksheet.cell(row, 1).value
            in {
                "Client",
                "Reporting Month",
                "Currency / Currency Status",
                "Total Markets Reviewed",
                "Markets with SLA Breach",
                "Successfully Calculated Breached Markets",
                "Ambiguous Breached Markets",
                "Unmatched Breached Markets",
            }
        }
        breach_header_row = next(
            row
            for row in range(1, worksheet.max_row + 1)
            if worksheet.cell(row, 1).value == "BREACHED MARKETS ONLY (SLA BREACH % > 0)"
        ) + 1
        breach_headers = [
            worksheet.cell(breach_header_row, column).value
            for column in range(1, worksheet.max_column + 1)
            if worksheet.cell(breach_header_row, column).value is not None
        ]
        confidence_index = breach_headers.index("Match Confidence") + 1
        breached_rows = []
        for row in range(breach_header_row + 1, worksheet.max_row + 1):
            if worksheet.cell(row, 1).value is not None:
                breached_rows.append(row)
        confidences = [
            worksheet.cell(row, confidence_index).value for row in breached_rows
        ]
        details_count = max(detail_sheet.max_row - 2, 0)
        detail_headers = [
            detail_sheet.cell(2, column).value
            for column in range(1, len(expected_detail_export.columns) + 1)
        ]
        detail_values_preserved = detail_headers == list(expected_detail_export.columns)
        run_id_column = (
            detail_headers.index("Processing Run ID") + 1
            if "Processing Run ID" in detail_headers
            else None
        )
        workbook_run_ids = set()
        for row_index, values in enumerate(
            expected_detail_export.itertuples(index=False, name=None),
            start=3,
        ):
            for column_index, expected_value in enumerate(values, start=1):
                if column_index == run_id_column:
                    continue
                detail_values_preserved = detail_values_preserved and _cell_matches(
                    detail_sheet.cell(row_index, column_index).value,
                    expected_value,
                )
            if run_id_column is not None:
                workbook_run_ids.add(detail_sheet.cell(row_index, run_id_column).value)
        detail_values_preserved = detail_values_preserved and (
            run_id_column is None
            or (len(workbook_run_ids) == 1 and all(bool(value) for value in workbook_run_ids))
        )
        expected_breached = sum(
            _is_positive(value) for value in expected_detail["SLA Breach %"].tolist()
        )
        expected_breached_rows = expected_detail.loc[
            expected_detail["SLA Breach %"].map(_is_positive)
        ]
        confidence_preserved = len(breached_rows) == len(expected_breached_rows)
        for row_index, (_, source_row) in zip(breached_rows, expected_breached_rows.iterrows()):
            confidence_preserved = confidence_preserved and (
                worksheet.cell(row_index, 1).value == source_row["DevOps Market"]
                and worksheet.cell(row_index, confidence_index).value
                == source_row["Match Confidence"]
                and worksheet.cell(
                    row_index,
                    breach_headers.index("Match Status") + 1,
                ).value
                == source_row["Match Status"]
            )
        check = {
            "first_sheet_is_executive_summary": workbook.sheetnames[0] == "Executive Summary",
            "breached_row_count_matches_control": len(breached_rows) == breached_count == expected_breached,
            "all_breached_rows_have_existing_confidence": (
                len(confidences) == breached_count
                and all(value in {"HIGH", "MEDIUM", "LOW"} for value in confidences)
            ),
            "match_confidence_and_status_unchanged_from_phase9": confidence_preserved,
            "total_market_count_matches_phase9": summary_rows.get("Total Markets Reviewed") == len(expected_detail),
            "detail_row_count_preserved": details_count == len(expected_detail),
            "detail_values_unchanged_from_phase9": detail_values_preserved,
            "ambiguous_unmatched_standard_credit_cells_blank": all(
                worksheet.cell(row, breach_headers.index("Calculated Credit Receivable") + 1).value is None
                for row in breached_rows
                if worksheet.cell(row, breach_headers.index("Match Status") + 1).value
                in {"AMBIGUOUS", "UNMATCHED"}
                and client != "CLIENT-004"
            ),
            "client004_low_unmatched_calculated_retained": (
                client != "CLIENT-004"
                or all(
                    worksheet.cell(row, breach_headers.index("Calculation Status") + 1).value
                    == "CALCULATED"
                    and worksheet.cell(row, confidence_index).value == "LOW"
                    and worksheet.cell(row, breach_headers.index("Match Status") + 1).value
                    == "UNMATCHED"
                    for row in breached_rows
                )
            ),
        }
        checks[client] = check
        workbook.close()
    return checks


def export_phase10_workbooks(
    detailed: pd.DataFrame,
    executive: pd.DataFrame,
    client_totals: pd.DataFrame,
    exceptions: pd.DataFrame,
    output_dir: str | Path,
    *,
    reporting_month: str = "2026-04",
) -> ExcelExportBundle:
    """Generate three client reports plus executive and exceptions workbooks."""
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    notes = _methodology_frame()
    expected: dict[str, dict[str, pd.DataFrame]] = {}
    paths: dict[str, Path] = {}

    for client in ("CLIENT-001", "CLIENT-002", "CLIENT-004"):
        filename = f"{client}_APR26_Credit_Report.xlsx"
        detail = _select_client_detail(detailed, client, reporting_month)
        sheets = {"Credit Detail": detail, "Notes & Methodology": notes}
        expected[filename] = sheets
        paths[filename] = directory / filename
        _write_workbook(
            paths[filename],
            sheets,
            {
                "Credit Detail": f"{client} - APRIL 2026 CREDIT DETAIL",
                "Notes & Methodology": "NOTES AND METHODOLOGY",
            },
        )

    executive_filename = "Executive_Report.xlsx"
    executive_sheets = {
        "Executive Summary": executive,
        "Client Totals": client_totals,
        "Notes & Methodology": notes,
    }
    expected[executive_filename] = executive_sheets
    paths[executive_filename] = directory / executive_filename
    _write_workbook(
        paths[executive_filename],
        executive_sheets,
        {
            "Executive Summary": "EXECUTIVE REPORT - AUTHORITATIVE MONTHS",
            "Client Totals": "CLIENT TOTALS BY CURRENCY",
            "Notes & Methodology": "NOTES AND METHODOLOGY",
        },
    )

    exceptions_filename = "Exceptions.xlsx"
    pending = exceptions.loc[exceptions["Exception Code"].eq("MONTH_CONFLICT")].copy()
    exception_sheets = {
        "Exceptions": exceptions,
        "Pending Month Conflict": pending.reset_index(drop=True),
        "Notes & Methodology": notes,
    }
    expected[exceptions_filename] = exception_sheets
    paths[exceptions_filename] = directory / exceptions_filename
    _write_workbook(
        paths[exceptions_filename],
        exception_sheets,
        {
            "Exceptions": "CREDIT REPORT EXCEPTIONS - ALL MONTHS",
            "Pending Month Conflict": "UNRESOLVED MAY26 - 66 ROWS",
            "Notes & Methodology": "NOTES AND METHODOLOGY",
        },
    )
    return ExcelExportBundle(paths=paths, expected_sheets=expected)


def _cell_matches(actual: Any, expected: Any) -> bool:
    expected_value = _excel_value(expected)
    if expected_value is None:
        return actual is None
    if isinstance(expected_value, Decimal):
        if actual is None or isinstance(actual, bool):
            return False
        try:
            return Decimal(str(actual)) == expected_value
        except Exception:
            return False
    if isinstance(expected_value, (int, float)) and not isinstance(expected_value, bool):
        if actual is None or isinstance(actual, bool):
            return False
        try:
            return Decimal(str(actual)) == Decimal(str(expected_value))
        except Exception:
            return False
    return actual == expected_value


def verify_phase10_workbooks(bundle: ExcelExportBundle) -> ExcelVerification:
    """Reopen workbooks and compare sheets, row counts, values, and formulas."""
    checks: dict[str, bool] = {}
    sheet_names: dict[str, tuple[str, ...]] = {}
    row_counts: dict[str, dict[str, int]] = {}
    open_check = True
    sheet_check = True
    row_check = True
    value_check = True
    formulas_absent = True

    for filename, path in bundle.paths.items():
        try:
            workbook = load_workbook(path, data_only=False)
        except Exception:
            open_check = False
            sheet_check = False
            row_check = False
            value_check = False
            formulas_absent = False
            continue
        actual_sheets = tuple(workbook.sheetnames)
        sheet_names[filename] = actual_sheets
        expected_frames = bundle.expected_sheets[filename]
        if set(actual_sheets) != set(expected_frames):
            sheet_check = False
        row_counts[filename] = {}
        for sheet_name, frame in expected_frames.items():
            if sheet_name not in workbook.sheetnames:
                sheet_check = False
                row_check = False
                value_check = False
                continue
            worksheet = workbook[sheet_name]
            data_rows = max(worksheet.max_row - 2, 0)
            row_counts[filename][sheet_name] = data_rows
            if data_rows != len(frame):
                row_check = False
            headers = [worksheet.cell(2, column).value for column in range(1, len(frame.columns) + 1)]
            if headers != list(frame.columns):
                value_check = False
            for row_index, values in enumerate(frame.itertuples(index=False, name=None), start=3):
                for column_index, expected_value in enumerate(values, start=1):
                    if not _cell_matches(
                        worksheet.cell(row_index, column_index).value,
                        expected_value,
                    ):
                        value_check = False
            if any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row):
                formulas_absent = False
        workbook.close()

    checks["files_open_successfully"] = open_check
    checks["expected_sheets_exist"] = sheet_check
    checks["row_counts_match_phase9"] = row_check
    checks["exported_values_match_phase9"] = value_check
    checks["no_excel_formulas_added"] = formulas_absent
    return ExcelVerification(checks, sheet_names, row_counts)


__all__ = [
    "ExcelExportBundle",
    "ExcelVerification",
    "add_client_executive_summaries",
    "export_phase10_workbooks",
    "verify_client_executive_summaries",
    "verify_phase10_workbooks",
]