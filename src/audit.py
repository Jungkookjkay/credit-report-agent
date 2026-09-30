"""Reusable Phase 12 control reconciliations for the assessment pipeline."""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import pandas as pd
from openpyxl import load_workbook

if __package__:
	from .models import CreditResult, MatchResult, SLARecord
	from .reporting import (
		CALCULATED_STATUSES,
		UNKNOWN_CURRENCY_BUCKET,
		build_exceptions_report,
	)
else:
	from models import CreditResult, MatchResult, SLARecord
	from reporting import CALCULATED_STATUSES, UNKNOWN_CURRENCY_BUCKET, build_exceptions_report


VALID_CONFIDENCES = {"HIGH", "MEDIUM", "LOW"}
PRIMARY_OUTCOMES = {"CALCULATED", "ZERO_CREDIT", "EXCEPTION", "INELIGIBLE"}
ASSESSMENT_CLIENT_ROW_COUNTS = {
	("CLIENT-001", "2026-04"): 56,
	("CLIENT-002", "2026-04"): 66,
	("CLIENT-004", "2026-04"): 56,
}
ASSESSMENT_BREACH_COUNTS = {
	"CLIENT-001": (7, 4, 2, 1),
	"CLIENT-002": (8, 3, 2, 3),
	"CLIENT-004": (4, 4, 0, 4),
}
ASSESSMENT_EXCEPTION_CODE_COUNTS = {
	("CLIENT-001", "AMBIGUOUS_MATCH"): 6,
	("CLIENT-001", "NON_UNIQUE_FINANCE_KEY"): 6,
	("CLIENT-001", "NO_FINANCE_MATCH"): 4,
	("CLIENT-002", "AMBIGUOUS_MATCH"): 15,
	("CLIENT-002", "MISSING_FINANCE_PARENT"): 1,
	("CLIENT-002", "MONTH_CONFLICT"): 66,
	("CLIENT-002", "NON_UNIQUE_FINANCE_KEY"): 4,
	("CLIENT-002", "NO_FINANCE_MATCH"): 17,
	("CLIENT-004", "CURRENCY_UNKNOWN"): 55,
	("CLIENT-004", "NO_FINANCE_MATCH"): 55,
}
ASSESSMENT_CREDIT_TOTALS = {
	("CLIENT-001", "EUR"): Decimal("1084.716168925"),
	("CLIENT-002", "USD"): Decimal("1240.742913354"),
	(
		"CLIENT-004",
		UNKNOWN_CURRENCY_BUCKET,
	): Decimal("21469.4194183928702104404702486177512721127279580188"),
}


@dataclass(frozen=True)
class ReconciliationControl:
	"""One auditable expected-versus-actual control outcome."""

	control: str
	expected: str
	actual: str
	status: str
	explanation: str


@dataclass(frozen=True)
class Phase12Reconciliation:
	"""Structured results for all Phase 12 controls."""

	controls: tuple[ReconciliationControl, ...]

	@property
	def passed_count(self) -> int:
		return sum(control.status == "PASS" for control in self.controls)

	@property
	def failed_count(self) -> int:
		return sum(control.status == "FAIL" for control in self.controls)

	@property
	def passed(self) -> bool:
		return self.failed_count == 0


def _decimal(value: Any) -> Decimal | None:
	if value is None or (isinstance(value, str) and not value.strip()):
		return None
	try:
		number = Decimal(str(value))
	except Exception:
		return None
	return number if number.is_finite() else None


def _source_key(record: Any) -> tuple[Any, Any, Any]:
	return (
		getattr(record, "source_file", None),
		getattr(record, "source_sheet", None),
		getattr(record, "source_row_number", None),
	)


def _is_breached(value: Any) -> bool:
	number = _decimal(value)
	return number is not None and number > 0


def _decimal_sum(values: Iterable[Any]) -> Decimal:
	numbers = [number for value in values if (number := _decimal(value)) is not None]
	with localcontext() as context:
		context.prec = 80
		return sum(numbers, Decimal("0"))


def _format_money_map(values: Mapping[tuple[str, str], Decimal]) -> str:
	return "; ".join(
		f"{client}/{currency}={format(amount, 'f')}"
		for (client, currency), amount in sorted(values.items())
	)


def _add_control(
	controls: list[ReconciliationControl],
	name: str,
	expected: Any,
	actual: Any,
	passed: bool,
	explanation: str,
) -> None:
	controls.append(
		ReconciliationControl(
			control=name,
			expected=str(expected),
			actual=str(actual),
			status="PASS" if passed else "FAIL",
			explanation=explanation,
		)
	)


def _result_exception_codes(result: CreditResult) -> tuple[str, ...]:
	codes = list(result.exception_codes)
	if result.exception_code and result.exception_code not in codes:
		codes.append(result.exception_code)
	for issue in result.source_exceptions:
		if issue.exception_code not in codes:
			codes.append(issue.exception_code)
	if result.month_conflict and "MONTH_CONFLICT" not in codes:
		codes.append("MONTH_CONFLICT")
	if any("currency is unknown" in note.lower() for note in result.audit_notes):
		if "CURRENCY_UNKNOWN" not in codes:
			codes.append("CURRENCY_UNKNOWN")
	return tuple(codes)


def _calculated(result: CreditResult) -> bool:
	return result.status in CALCULATED_STATUSES and result.calculated_credit is not None


def _result_money_by_currency(results: Sequence[CreditResult]) -> dict[tuple[str, str], Decimal]:
	grouped: dict[tuple[str, str], list[Any]] = defaultdict(list)
	for result in results:
		if result.reporting_month != "2026-04" or not _calculated(result):
			continue
		currency = (
			UNKNOWN_CURRENCY_BUCKET
			if result.currency is None or not str(result.currency).strip()
			else str(result.currency)
		)
		grouped[(str(result.client), currency)].append(result.calculated_credit)
	return {key: _decimal_sum(values) for key, values in grouped.items()}


def _frame_money_by_currency(frame: pd.DataFrame) -> dict[tuple[str, str], Decimal]:
	if frame.empty:
		return {}
	grouped: dict[tuple[str, str], list[Any]] = defaultdict(list)
	for _, row in frame.iterrows():
		if row.get("Reporting Month") != "2026-04":
			continue
		amount = _decimal(row.get("Total Calculated Credit"))
		if amount is None:
			continue
		grouped[(str(row.get("Client")), str(row.get("Currency")))].append(amount)
	return {key: _decimal_sum(values) for key, values in grouped.items()}


def _reference_code_counter(frame: pd.DataFrame) -> Counter:
	return Counter(
		(
			row.get("DevOps Source File"),
			row.get("DevOps Source Sheet"),
			row.get("DevOps Source Row"),
			row.get("Client"),
			row.get("Exception Code"),
		)
		for _, row in frame.iterrows()
	)


def _load_excel_frame(path: Path, sheet_name: str, header_row: int = 2) -> pd.DataFrame:
	workbook = load_workbook(path, data_only=False, read_only=True)
	try:
		worksheet = workbook[sheet_name]
		rows = worksheet.iter_rows(min_row=header_row, values_only=True)
		headers = next(rows)
		records = [dict(zip(headers, row)) for row in rows]
		return pd.DataFrame(records)
	finally:
		workbook.close()


def _summary_values(worksheet: Any) -> dict[str, Any]:
	return {
		worksheet.cell(row, 1).value: worksheet.cell(row, 2).value
		for row in range(4, worksheet.max_row + 1)
	}


def _client_workbook_checks(
	output_dir: Path,
	detailed: pd.DataFrame,
	results_by_key: Mapping[tuple[Any, Any, Any], CreditResult],
) -> tuple[dict[str, bool], dict[str, Any]]:
	checks: dict[str, bool] = {}
	details: dict[str, Any] = {}
	breached_table_confidence_ok = True
	detail_lineage_ok = True
	detail_counts_ok = True
	summary_counts_ok = True
	summary_controls_ok = True
	detail_counts: dict[str, int] = {}
	summary_breached_counts: dict[str, int] = {}
	confidence_mismatches: list[str] = []
	results_by_market = {
		(result.client, result.devops_market): result
		for result in results_by_key.values()
		if result.reporting_month == "2026-04"
		and result.reporting_month_status == "RESOLVED"
	}

	for client, expected_detail_count in (
		("CLIENT-001", 56),
		("CLIENT-002", 66),
		("CLIENT-004", 56),
	):
		path = output_dir / f"{client}_APR26_Credit_Report.xlsx"
		if not path.exists():
			checks[f"{client}_workbook_exists"] = False
			details[client] = "missing workbook"
			continue
		workbook = load_workbook(path, data_only=False, read_only=True)
		try:
			if workbook.sheetnames != [
				"Executive Summary",
				"Credit Detail",
				"Notes & Methodology",
			]:
				checks[f"{client}_sheet_structure"] = False
			else:
				checks[f"{client}_sheet_structure"] = True
			summary = workbook["Executive Summary"]
			detail = workbook["Credit Detail"]
			detail_headers = [cell.value for cell in next(detail.iter_rows(min_row=2, max_row=2))]
			detail_column = {name: index + 1 for index, name in enumerate(detail_headers)}
			detail_records = list(detail.iter_rows(min_row=3, values_only=True))
			detail_records = [row for row in detail_records if any(value is not None for value in row)]
			detail_counts[client] = len(detail_records)
			details[f"{client}_detail_count"] = len(detail_records)
			detail_counts_ok = detail_counts_ok and len(detail_records) == expected_detail_count

			for row in detail_records:
				source_key = (
					row[detail_column["DevOps Source File"] - 1],
					row[detail_column["DevOps Source Sheet"] - 1],
					row[detail_column["DevOps Source Row"] - 1],
				)
				result = results_by_key.get(source_key)
				if result is None:
					detail_lineage_ok = False
					continue
				if row[detail_column["Confidence"] - 1] != result.match_confidence:
					breached_table_confidence_ok = False
					confidence_mismatches.append(
						f"{client} detail {source_key} Confidence: "
						f"expected={result.match_confidence!r}, actual={row[detail_column['Confidence'] - 1]!r}"
					)
				if row[detail_column["Match Status"] - 1] != result.match_status:
					breached_table_confidence_ok = False
					confidence_mismatches.append(
						f"{client} detail {source_key} Match Status: "
						f"expected={result.match_status!r}, actual={row[detail_column['Match Status'] - 1]!r}"
					)
				if row[detail_column["Finance Source File"] - 1] != result.finance_source_file:
					detail_lineage_ok = False
				if row[detail_column["Finance Source Sheet"] - 1] != result.finance_source_sheet:
					detail_lineage_ok = False
				if row[detail_column["Finance Source Row"] - 1] != result.finance_source_row_number:
					detail_lineage_ok = False

			summary_values = _summary_values(summary)
			expected_client_rows = sum(
				value["Client"] == client
				and value["Reporting Month"] == "2026-04"
				and value["Reporting Month Status"] == "RESOLVED"
				for _, value in detailed.iterrows()
			)
			if summary_values.get("Total Markets Reviewed") != expected_client_rows:
				summary_counts_ok = False
			breached_header_row = next(
				row_number
				for row_number in range(1, summary.max_row + 1)
				if summary.cell(row_number, 1).value
				== "BREACHED MARKETS ONLY (SLA BREACH % > 0)"
			) + 1
			headers = {
				summary.cell(breached_header_row, col).value: col
				for col in range(1, summary.max_column + 1)
				if summary.cell(breached_header_row, col).value is not None
			}
			breached_rows = []
			for row_number in range(breached_header_row + 1, summary.max_row + 1):
				if summary.cell(row_number, 1).value is not None:
					breached_rows.append(row_number)
			expected_breached = ASSESSMENT_BREACH_COUNTS[client][0]
			summary_breached_counts[client] = len(breached_rows)
			summary_counts_ok = summary_counts_ok and len(breached_rows) == expected_breached
			for row_number in breached_rows:
				if summary.cell(row_number, headers["Match Confidence"]).value not in VALID_CONFIDENCES:
					breached_table_confidence_ok = False
				result = results_by_market.get(
					(client, summary.cell(row_number, headers["DevOps Market"]).value)
				)
				if result is None:
					breached_table_confidence_ok = False
				elif (
					summary.cell(row_number, headers["Match Confidence"]).value
					!= result.match_confidence
					or summary.cell(row_number, headers["Match Status"]).value
					!= result.match_status
				):
					breached_table_confidence_ok = False
					confidence_mismatches.append(
						f"{client} summary market={result.devops_market!r}: "
						f"expected=({result.match_confidence!r}, {result.match_status!r}), "
						f"actual=({summary.cell(row_number, headers['Match Confidence']).value!r}, "
						f"{summary.cell(row_number, headers['Match Status']).value!r})"
					)

			expected_rule_counts = ASSESSMENT_BREACH_COUNTS[client]
			actual_summary_counts = (
				summary_values.get("Markets with SLA Breach"),
				summary_values.get("Successfully Calculated Breached Markets"),
				summary_values.get("Ambiguous Breached Markets"),
				summary_values.get("Unmatched Breached Markets"),
			)
			summary_controls_ok = summary_controls_ok and (
				actual_summary_counts == expected_rule_counts
			)
			details[f"{client}_breached_table_count"] = len(breached_rows)
		finally:
			workbook.close()

	checks["authoritative_client_detail_row_counts"] = detail_counts_ok
	details["client_detail_row_counts"] = detail_counts
	checks["client_summary_breached_counts"] = summary_counts_ok
	details["client_summary_breached_counts"] = summary_breached_counts
	details["confidence_mismatches"] = confidence_mismatches
	checks["client_summary_operational_counts"] = summary_controls_ok
	checks["confidence_and_match_status_preserved_in_excel"] = breached_table_confidence_ok
	checks["finance_and_devops_lineage_preserved_in_excel_detail"] = detail_lineage_ok
	return checks, details


def run_phase12_reconciliations(
	sla_records: Sequence[SLARecord],
	match_results: Sequence[MatchResult],
	credit_results: Sequence[CreditResult],
	detailed: pd.DataFrame,
	client_totals: pd.DataFrame,
	executive: pd.DataFrame,
	exceptions: pd.DataFrame,
	*,
	output_dir: str | Path,
) -> Phase12Reconciliation:
	"""Run explicit population, month, financial, exception, confidence, and lineage controls."""
	controls: list[ReconciliationControl] = []
	output_path = Path(output_dir)
	sla_keys = [_source_key(record) for record in sla_records]
	result_keys = [_source_key(result) for result in credit_results]
	outcome_counts = Counter(result.status for result in credit_results)
	valid_outcomes = all(result.status in PRIMARY_OUTCOMES for result in credit_results)
	one_to_one = (
		len(sla_records) == len(credit_results) == len(match_results)
		and len(set(sla_keys)) == len(sla_keys)
		and Counter(sla_keys) == Counter(result_keys)
		and valid_outcomes
	)
	outcome_detail = ", ".join(
		f"{name}={outcome_counts.get(name, 0)}"
		for name in sorted(PRIMARY_OUTCOMES)
	)
	_add_control(
		controls,
		"SLA population has exactly one primary outcome",
		"244 SLA rows; 244 aligned MatchResults; 244 unique CreditResults",
		f"SLA={len(sla_records)}; MatchResult={len(match_results)}; CreditResult={len(credit_results)}; outcomes: {outcome_detail}",
		one_to_one,
		"Exception-code multiplicity is not counted as additional SLA outcomes; source file/sheet/row identities must align one-to-one.",
	)

	month_counts = Counter(
		(result.client, result.reporting_month, result.reporting_month_status)
		for result in credit_results
	)
	expected_month_counts = {
		("CLIENT-001", "2026-04", "RESOLVED"): 56,
		("CLIENT-002", "2026-04", "RESOLVED"): 66,
		("CLIENT-004", "2026-04", "RESOLVED"): 56,
	}
	month_counts_ok = all(month_counts.get(key, 0) == value for key, value in expected_month_counts.items())
	unresolved_may = [
		result for result in credit_results
		if result.client == "CLIENT-002"
		and result.source_file
		and "MAY26" in result.source_file.upper()
	]
	month_counts_ok = month_counts_ok and len(unresolved_may) == 66 and all(
		result.reporting_month is None
		and result.reporting_month_status == "UNRESOLVED"
		and result.status == "INELIGIBLE"
		and "MONTH_CONFLICT" in _result_exception_codes(result)
		and result.calculated_credit is None
		for result in unresolved_may
	)
	unresolved_count = sum(result.reporting_month_status != "RESOLVED" for result in credit_results)
	month_counts_ok = month_counts_ok and unresolved_count == 66
	month_actual = "; ".join(
		f"{client}/{month or 'UNRESOLVED'}={count}"
		for (client, month, _status), count in sorted(month_counts.items(), key=lambda item: str(item[0]))
	)
	_add_control(
		controls,
		"Client/month and unresolved MAY26 populations",
		"CLIENT-001 APR=56; CLIENT-002 APR=66; CLIENT-004 APR=56; CLIENT-002 MAY unresolved=66",
		f"{month_actual}; unresolved_total={unresolved_count}",
		month_counts_ok,
		"Only resolved 2026-04 rows enter authoritative monthly reporting; MAY26 filename/content conflict stays unresolved and excluded.",
	)

	breached_results = [result for result in credit_results if _is_breached(result.sla_breach_pct)]
	eligible_breached = [
		result for result in breached_results
		if result.match_status != "INELIGIBLE"
	]
	pre_match_excluded = [
		result for result in breached_results
		if result.match_status == "INELIGIBLE"
	]
	null_confidence_eligible = sum(
		result.match_confidence not in VALID_CONFIDENCES
		for result in eligible_breached
	)
	excluded_are_explicit = len(pre_match_excluded) == 3 and all(
		result.client == "CLIENT-002"
		and result.status == "INELIGIBLE"
		and result.exception_code == "MONTH_CONFLICT"
		and result.match_confidence is None
		and result.reporting_month_status == "UNRESOLVED"
		for result in pre_match_excluded
	)
	confidence_population_ok = (
		len(breached_results) == 22
		and len(eligible_breached) == 19
		and len(pre_match_excluded) == 3
		and null_confidence_eligible == 0
		and excluded_are_explicit
	)
	_add_control(
		controls,
		"Breached population confidence eligibility",
		"22 breached = 19 eligible with HIGH/MEDIUM/LOW + 3 pre-match MONTH_CONFLICT/INELIGIBLE with null confidence",
		f"breached={len(breached_results)}; eligible={len(eligible_breached)}; eligible_null_confidence={null_confidence_eligible}; excluded_null_confidence={len(pre_match_excluded)}",
		confidence_population_ok,
		"Confidence applies only to eligible records where Finance matching was attempted. The three MAY26 records were stopped before matching, retain null confidence, and are explicitly explained as MONTH_CONFLICT/INELIGIBLE.",
	)

	client_breach_actual: dict[str, tuple[int, int, int, int]] = {}
	client_breach_ok = True
	for client, expected_counts in ASSESSMENT_BREACH_COUNTS.items():
		rows = [
			result for result in credit_results
			if result.client == client and result.reporting_month == "2026-04"
		]
		breached = [result for result in rows if _is_breached(result.sla_breach_pct)]
		calculated_breached = [
			result for result in breached if _calculated(result)
		]
		ambiguous = sum(result.match_status == "AMBIGUOUS" for result in breached)
		unmatched = sum(result.match_status == "UNMATCHED" for result in breached)
		client_breach_actual[client] = (
			len(breached), len(calculated_breached), ambiguous, unmatched
		)
		client_breach_ok = client_breach_ok and client_breach_actual[client] == expected_counts
	client_breach_ok = client_breach_ok and all(
		result.match_status == "UNMATCHED"
		and result.match_confidence == "LOW"
		and result.status == "CALCULATED"
		and result.calculation_rule == "CLIENT_004_SPECIAL"
		for result in credit_results
		if result.client == "CLIENT-004"
		and result.reporting_month == "2026-04"
		and _is_breached(result.sla_breach_pct)
	)
	_add_control(
		controls,
		"Breached-market outcomes by client",
		"CLIENT-001=7/4/2/1; CLIENT-002=8/3/2/3; CLIENT-004=4/4/0/4 (breached/calculated/ambiguous/unmatched)",
		"; ".join(
			f"{client}=" + "/".join(str(value) for value in counts)
			for client, counts in sorted(client_breach_actual.items())
		),
		client_breach_ok,
		"CLIENT-004 LOW/UNMATCHED remains a successful special-rule calculation; Finance-link status and calculation status are independently represented.",
	)

	detailed_source_counts = (
		detailed.groupby("DevOps Source File", dropna=False).size().to_dict()
		if not detailed.empty
		else {}
	)
	detailed_count_ok = len(detailed) == 244 and sum(detailed_source_counts.values()) == 244
	authoritative_detail_count = int(detailed["Reporting Month Status"].eq("RESOLVED").sum())
	detailed_count_ok = detailed_count_ok and authoritative_detail_count == 178
	_add_control(
		controls,
		"Phase 9 detailed report population",
		"244 total detail rows; 178 authoritative April rows; per-source 56/66/66/56",
		f"rows={len(detailed)}; authoritative_April={authoritative_detail_count}; by_source={detailed_source_counts}",
		detailed_count_ok,
		"The 66 unresolved MAY26 rows remain present in detail but are not authoritative April detail or financial totals.",
	)

	results_by_key = {_source_key(result): result for result in credit_results}
	workbook_checks, workbook_details = _client_workbook_checks(
		output_path, detailed, results_by_key
	)
	_add_control(
		controls,
		"Excel client detail and breached-summary reconciliation",
		"Detail rows=56/66/56; summary breached rows=7/8/4; summary metrics/confidence/lineage unchanged",
		f"checks={workbook_checks}; details={workbook_details}",
		all(workbook_checks.values()),
		"Existing workbooks are opened read-only. Summary counts, confidence/status, and detail lineage are checked against Phase 9/CreditResult data.",
	)

	result_money = _result_money_by_currency(credit_results)
	detail_money = _frame_money_by_currency(
		detailed.loc[
			detailed["Reporting Month"].eq("2026-04")
			& detailed["Calculation Status"].isin(CALCULATED_STATUSES)
			& detailed["Calculated Credit"].notna(),
			["Reporting Month", "Client", "Reporting Currency Bucket", "Calculated Credit"],
		].rename(columns={"Reporting Currency Bucket": "Currency", "Calculated Credit": "Total Calculated Credit"})
	)
	executive_money = _frame_money_by_currency(executive)
	financial_actual = _format_money_map(result_money)
	financial_totals_ok = all(result_money.get(key) == amount for key, amount in ASSESSMENT_CREDIT_TOTALS.items())
	financial_totals_ok = financial_totals_ok and detail_money == result_money and executive_money == result_money
	_add_control(
		controls,
		"Exact calculated detail-to-executive totals",
		_format_money_map(ASSESSMENT_CREDIT_TOTALS),
		f"CreditResult={financial_actual}; detail_matches={detail_money == result_money}; executive_matches={executive_money == result_money}",
		financial_totals_ok,
		"Calculated Decimal values only are summed by authoritative month/client/currency; display rounding is not used in this reconciliation.",
	)

	executive_keys = [
		(row["Reporting Month"], row["Client"], row["Currency"])
		for _, row in executive.iterrows()
	]
	currency_separation_ok = len(executive_keys) == len(set(executive_keys))
	currency_separation_ok = currency_separation_ok and (
		("CLIENT-004", UNKNOWN_CURRENCY_BUCKET) in result_money
		and all(
			result.currency is None
			for result in credit_results
			if result.client == "CLIENT-004"
			and result.calculation_rule == "CLIENT_004_SPECIAL"
			and result.match_status == "UNMATCHED"
			and _calculated(result)
		)
	)
	null_credit_count = sum(result.calculated_credit is None for result in credit_results)
	zero_credit_count = sum(
		_calculated(result) and _decimal(result.calculated_credit) == 0
		for result in credit_results
	)
	currency_separation_ok = currency_separation_ok and zero_credit_count > 0 and null_credit_count > 0
	_add_control(
		controls,
		"Currency separation and null/zero semantics",
		"No mixed-currency executive keys; CLIENT-004 unmatched specials stay unknown; null excluded and zero retained",
		f"unique_exec_groups={len(executive_keys) == len(set(executive_keys))}; C4_unknown_total={result_money.get(('CLIENT-004', UNKNOWN_CURRENCY_BUCKET))}; null_credits={null_credit_count}; valid_zero_credits={zero_credit_count}",
		currency_separation_ok,
		"Unknown-currency CLIENT-004 amounts are not converted or combined with USD. Null credits are not substituted with zero; valid zero-credit results remain in calculated populations.",
	)

	expected_code_counter = Counter(ASSESSMENT_EXCEPTION_CODE_COUNTS)
	actual_code_counter = Counter(
		(str(row["Client"]), str(row["Exception Code"]))
		for _, row in exceptions.iterrows()
	)
	exception_code_ok = actual_code_counter == expected_code_counter
	expected_exception_output = build_exceptions_report(credit_results)
	expected_ref_codes = _reference_code_counter(expected_exception_output)
	actual_ref_codes = _reference_code_counter(exceptions)
	exception_code_ok = exception_code_ok and expected_ref_codes == actual_ref_codes
	_add_control(
		controls,
		"Exception-code occurrences match processing results",
		f"approved code occurrences={dict(sorted(expected_code_counter.items()))}",
		f"code_occurrences={len(exceptions)}; counts={dict(sorted(actual_code_counter.items()))}",
		exception_code_ok,
		"Counts are exception-code occurrences. A single SLA record may carry multiple codes and is not counted as multiple primary outcomes.",
	)

	expected_unique_refs = {
		(row["DevOps Source File"], row["DevOps Source Sheet"], row["DevOps Source Row"])
		for _, row in expected_exception_output.iterrows()
	}
	actual_unique_refs = {
		(row["DevOps Source File"], row["DevOps Source Sheet"], row["DevOps Source Row"])
		for _, row in exceptions.iterrows()
	}
	_add_control(
		controls,
		"Unique affected SLA exception population",
		str(len(expected_unique_refs)),
		f"expected_unique={len(expected_unique_refs)}; output_unique={len(actual_unique_refs)}; code_rows={len(exceptions)}",
		expected_unique_refs == actual_unique_refs,
		"Unique source file/sheet/row references are reconciled separately from the 229 exception-code rows.",
	)

	month_conflict_excel_ok = False
	excel_exception_counts: Counter = Counter()
	try:
		exception_path = output_path / "Exceptions.xlsx"
		exception_frame = _load_excel_frame(exception_path, "Exceptions")
		excel_exception_counts = Counter(
			(str(row["Client"]), str(row["Exception Code"]))
			for _, row in exception_frame.iterrows()
		)
		pending_frame = _load_excel_frame(exception_path, "Pending Month Conflict")
		pending_ok = (
			len(pending_frame) == 66
			and pending_frame["Exception Code"].eq("MONTH_CONFLICT").all()
			and pending_frame["Reporting Month Status"].eq("UNRESOLVED").all()
			and pending_frame["Reporting Month"].isna().all()
		)
		month_conflict_excel_ok = (
			pending_ok
			and excel_exception_counts == actual_code_counter
			and excel_exception_counts[("CLIENT-002", "MONTH_CONFLICT")] == 66
		)
	except Exception:
		pending_ok = False
	_add_control(
		controls,
		"Excel exception output and 66 pending month conflicts",
		"229 exception-code rows; 66 unresolved MONTH_CONFLICT rows",
		f"excel_exception_code_rows={sum(excel_exception_counts.values())}; MAY26_pending={66 if pending_ok else 'not reconciled'}",
		month_conflict_excel_ok,
		"The Exceptions workbook is read-only checked; all MAY26 conflict rows remain unresolved and outside authoritative totals.",
	)

	eligible_confidence_ok = all(
		result.match_confidence in VALID_CONFIDENCES
		for result in eligible_breached
	)
	excluded_confidence_ok = excluded_are_explicit and all(
		result.match_confidence is None for result in pre_match_excluded
	)
	chain_ok = len(match_results) == len(credit_results) and len(detailed) == len(credit_results)
	for index, (match, result) in enumerate(zip(match_results, credit_results)):
		chain_ok = chain_ok and (
			match.status == result.match_status
			and match.confidence == result.match_confidence
			and match.match_method == result.match_method
			and detailed.iloc[index]["Match Status"] == match.status
			and detailed.iloc[index]["Match Confidence"] == match.confidence
		)
	_add_control(
		controls,
		"Confidence propagation and eligibility",
		"19 eligible breached records retain source confidence; 3 excluded breaches retain null confidence with MONTH_CONFLICT",
		f"eligible_valid={sum(result.match_confidence in VALID_CONFIDENCES for result in eligible_breached)}/{len(eligible_breached)}; eligible_null={null_confidence_eligible}; excluded_null={sum(result.match_confidence is None for result in pre_match_excluded)}/{len(pre_match_excluded)}; MatchResult_to_CreditResult_to_detail={chain_ok}",
		eligible_confidence_ok and excluded_confidence_ok and chain_ok,
		"Confidence is a Finance-match reliability signal. The three MONTH_CONFLICT/INELIGIBLE records were excluded before matching and intentionally retain None, not LOW.",
	)

	lineage_ok = len(sla_records) == len(credit_results)
	for sla, match, result in zip(sla_records, match_results, credit_results):
		lineage_ok = lineage_ok and (
			_source_key(sla) == _source_key(result)
			and result.devops_market == sla.raw_market_code
		)
		if _calculated(result):
			lineage_ok = lineage_ok and all(_source_key(result))
		if result.calculation_rule == "STANDARD_CAPPED" and _calculated(result):
			finance_record = match.matched_finance_record
			if finance_record is None and len(match.matched_finance_records) == 1:
				finance_record = match.matched_finance_records[0]
			lineage_ok = lineage_ok and (
				match.status == "MATCHED"
				and match.confidence == "HIGH"
				and finance_record is not None
				and result.finance_source_file == finance_record.source_file
				and result.finance_source_sheet == finance_record.source_sheet
				and result.finance_source_row_number == finance_record.source_row_number
				and _decimal(result.base_credit) == _decimal(finance_record.base_credit)
				and result.currency == finance_record.currency
			)
		if (
			result.client == "CLIENT-004"
			and result.calculation_rule == "CLIENT_004_SPECIAL"
			and match.status == "UNMATCHED"
			and _calculated(result)
		):
			lineage_ok = lineage_ok and all(
				value is None
				for value in (
					result.finance_source_file,
					result.finance_source_sheet,
					result.finance_source_row_number,
					result.matched_finance_market,
					result.customer_wd,
					result.base_credit,
					result.currency,
				)
			)
	detail_lineage_columns = {
		"DevOps Source File",
		"DevOps Source Sheet",
		"DevOps Source Row",
	}
	lineage_ok = lineage_ok and detail_lineage_columns.issubset(detailed.columns)
	for index, result in enumerate(credit_results):
		row = detailed.iloc[index]
		lineage_ok = lineage_ok and (
			row["DevOps Source File"] == result.source_file
			and row["DevOps Source Sheet"] == result.source_sheet
			and row["DevOps Source Row"] == result.source_row_number
		)
	lineage_ok = lineage_ok and workbook_checks.get("finance_and_devops_lineage_preserved_in_excel_detail", False)
	_add_control(
		controls,
		"Calculated credit DevOps/Finance lineage",
		"Every calculated result has DevOps lineage; standard results also match selected Finance lineage/Base Credit/currency; unmatched CLIENT-004 has no invented Finance lineage",
		f"all_calculated_and_standard_finance_lineage_valid={lineage_ok}; workbook_detail_lineage={workbook_checks.get('finance_and_devops_lineage_preserved_in_excel_detail', False)}",
		lineage_ok,
		"CLIENT-004 special calculations retain their DevOps origin. Finance source fields are required only for calculations that actually used a matched Finance record.",
	)

	excel_financial_ok = True
	excel_financial_actual: dict[tuple[str, str], Decimal] = {}
	try:
		executive_path = output_path / "Executive_Report.xlsx"
		excel_executive = _load_excel_frame(executive_path, "Executive Summary")
		for key in ASSESSMENT_CREDIT_TOTALS:
			client, currency = key
			rows = excel_executive.loc[
				excel_executive["Client"].eq(client)
				& excel_executive["Currency"].eq(currency)
			]
			if len(rows) != 1:
				excel_financial_ok = False
				continue
			amount = _decimal(rows.iloc[0]["Total Calculated Credit"])
			excel_financial_actual[key] = amount if amount is not None else Decimal("0")
			excel_financial_ok = excel_financial_ok and amount == ASSESSMENT_CREDIT_TOTALS[key]
		if "2026-05" in set(excel_executive["Reporting Month"].dropna()):
			excel_financial_ok = False
		for client in ("CLIENT-001", "CLIENT-002", "CLIENT-004"):
			workbook = load_workbook(
				output_path / f"{client}_APR26_Credit_Report.xlsx",
				data_only=False,
				read_only=True,
			)
			try:
				summary = workbook["Executive Summary"]
				values = _summary_values(summary)
				total = _decimal(values.get("Authoritative Total (Exact)"))
				expected_key = {
					"CLIENT-001": ("CLIENT-001", "EUR"),
					"CLIENT-002": ("CLIENT-002", "USD"),
					"CLIENT-004": ("CLIENT-004", UNKNOWN_CURRENCY_BUCKET),
				}[client]
				excel_financial_ok = excel_financial_ok and total == ASSESSMENT_CREDIT_TOTALS[expected_key]
			finally:
				workbook.close()
	except Exception:
		excel_financial_ok = False
	_add_control(
		controls,
		"Excel display totals preserve authoritative totals",
		_format_money_map(ASSESSMENT_CREDIT_TOTALS),
		_format_money_map(excel_financial_actual),
		excel_financial_ok,
		"Client summary display rounding is checked separately from its exact-value field and Executive Report values; the unknown-currency total remains text where required for precision.",
	)

	return Phase12Reconciliation(tuple(controls))


__all__ = [
	"Phase12Reconciliation",
	"ReconciliationControl",
	"run_phase12_reconciliations",
]
