"""Focused tests for Phase 10 Excel exports and reopen verification."""

from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook

from src.excel_export import (
    add_client_executive_summaries,
    export_phase10_workbooks,
    verify_phase10_workbooks,
)
from src.models import CreditResult, FinanceRecord
from src.reporting import (
    UNKNOWN_CURRENCY_BUCKET,
    build_client_totals,
    build_detailed_report,
    build_exceptions_report,
    build_executive_summary,
)


def _result(
    *,
    client: str,
    row: int,
    credit: Decimal | None,
    currency: str | None,
    status: str = "CALCULATED",
    month: str | None = "2026-04",
    match_status: str = "MATCHED",
    confidence: str | None = "HIGH",
    rule: str = "STANDARD_CAPPED",
    exception_codes: tuple[str, ...] = (),
    candidate_finance_records: tuple[FinanceRecord, ...] = (),
) -> CreditResult:
    is_matched = match_status == "MATCHED"
    return CreditResult(
        client=client,
        devops_market=f"market-{row}",
        normalized_devops_market=f"MARKET-{row}",
        matched_finance_market="FIN-CODE" if is_matched else None,
        customer_wd=f"{client}-WD-01" if is_matched else None,
        match_status=match_status,
        match_confidence=confidence,
        match_method="EXACT_CODE_MATCH" if is_matched else None,
        candidate_count=len(candidate_finance_records),
        candidate_finance_records=candidate_finance_records,
        sla_breach_pct=Decimal("0.05"),
        applied_rate=Decimal("0.05") if rule == "STANDARD_CAPPED" else None,
        base_credit=Decimal("1000") if is_matched else None,
        currency=currency,
        calculation_rule=rule,
        calculated_credit=credit,
        status=status,
        exception_code=exception_codes[0] if status in {"EXCEPTION", "INELIGIBLE"} and exception_codes else None,
        exception_codes=exception_codes,
        audit_notes=("Currency is unknown: no unique Finance record matched.",)
        if rule == "CLIENT_004_SPECIAL" and currency is None
        else (),
        reporting_month=month,
        reporting_month_status="RESOLVED" if month else "UNRESOLVED",
        filename_month_hint="APR26" if month else "MAY26",
        workbook_month_hint="April 2026",
        month_conflict=month is None,
        source_file=f"{client}-{row}.xlsx",
        source_sheet="SLA_Breaches",
        source_row_number=row,
        finance_source_file="finance.xlsx" if is_matched else None,
        finance_source_sheet="Finance_Fines" if is_matched else None,
        finance_source_row_number=10 if is_matched else None,
        processing_run_id="run-test",
    )


def test_excel_exports_reopen_with_expected_sheets_rows_and_values(tmp_path) -> None:
    long_credit = Decimal(
        "21469.4194183928702104404702486177512721127279580188"
    )
    candidates = (
        FinanceRecord(customer_wd="CLIENT-001-WD-01", raw_market_code="MX", base_credit=100, currency="EUR"),
        FinanceRecord(customer_wd="CLIENT-001-WD-02", raw_market_code="MX", base_credit=200, currency="EUR"),
    )
    results = [
        _result(
            client="CLIENT-001",
            row=1,
            credit=Decimal("1084.716168925"),
            currency="EUR",
        ),
        _result(
            client="CLIENT-001",
            row=2,
            credit=None,
            currency=None,
            status="EXCEPTION",
            match_status="AMBIGUOUS",
            confidence="LOW",
            exception_codes=("AMBIGUOUS_MATCH",),
            candidate_finance_records=candidates,
        ),
        _result(
            client="CLIENT-002",
            row=3,
            credit=Decimal("1240.742913354"),
            currency="USD",
        ),
        _result(
            client="CLIENT-002",
            row=4,
            credit=None,
            currency=None,
            status="INELIGIBLE",
            month=None,
            match_status="INELIGIBLE",
            confidence=None,
            exception_codes=("MONTH_CONFLICT",),
        ),
        _result(
            client="CLIENT-004",
            row=5,
            credit=long_credit,
            currency=None,
            rule="CLIENT_004_SPECIAL",
            match_status="UNMATCHED",
            confidence="LOW",
            exception_codes=("NO_FINANCE_MATCH",),
        ),
        _result(
            client="CLIENT-004",
            row=6,
            credit=Decimal("0"),
            currency="USD",
            status="ZERO_CREDIT",
            rule="CLIENT_004_SPECIAL",
        ),
    ]
    detailed = build_detailed_report(results)
    client_totals = build_client_totals(results)
    executive = build_executive_summary(results)
    exceptions = build_exceptions_report(results)

    bundle = export_phase10_workbooks(
        detailed,
        executive,
        client_totals,
        exceptions,
        tmp_path,
    )
    verification = verify_phase10_workbooks(bundle)

    assert verification.passed, verification.checks
    assert set(bundle.paths) == {
        "CLIENT-001_APR26_Credit_Report.xlsx",
        "CLIENT-002_APR26_Credit_Report.xlsx",
        "CLIENT-004_APR26_Credit_Report.xlsx",
        "Executive_Report.xlsx",
        "Exceptions.xlsx",
    }
    assert verification.row_counts["CLIENT-001_APR26_Credit_Report.xlsx"]["Credit Detail"] == 2
    assert verification.row_counts["CLIENT-002_APR26_Credit_Report.xlsx"]["Credit Detail"] == 1
    assert verification.row_counts["Exceptions.xlsx"]["Pending Month Conflict"] == 1
    assert verification.sheet_names["Executive_Report.xlsx"] == (
        "Executive Summary",
        "Client Totals",
        "Notes & Methodology",
    )

    detail_book = load_workbook(bundle.paths["CLIENT-001_APR26_Credit_Report.xlsx"])
    detail_sheet = detail_book["Credit Detail"]
    headers = {cell.value: cell.column for cell in detail_sheet[2]}
    assert detail_sheet.freeze_panes == "A3"
    assert detail_sheet.auto_filter.ref
    assert detail_sheet.cell(3, headers["Calculated Credit"]).value == 1084.716168925
    assert detail_sheet.cell(4, headers["Calculated Credit"]).value is None
    assert detail_sheet.cell(3, headers["SLA Breach %"]).number_format == "0.00%"
    detail_book.close()

    executive_book = load_workbook(bundle.paths["Executive_Report.xlsx"])
    executive_sheet = executive_book["Executive Summary"]
    executive_headers = {cell.value: cell.column for cell in executive_sheet[2]}
    client_column = executive_headers["Client"]
    currency_column = executive_headers["Currency"]
    amount_column = executive_headers["Total Calculated Credit"]
    rows_by_currency = {
        (
            executive_sheet.cell(row, client_column).value,
            executive_sheet.cell(row, currency_column).value,
        ): executive_sheet.cell(row, amount_column).value
        for row in range(3, executive_sheet.max_row + 1)
        if executive_sheet.cell(row, currency_column).value in {"USD", UNKNOWN_CURRENCY_BUCKET}
    }
    assert rows_by_currency[("CLIENT-004", UNKNOWN_CURRENCY_BUCKET)] == str(long_credit)
    assert rows_by_currency[("CLIENT-002", "USD")] == 1240.742913354
    assert rows_by_currency[("CLIENT-004", "USD")] == 0
    executive_book.close()

    exceptions_book = load_workbook(bundle.paths["Exceptions.xlsx"])
    pending_sheet = exceptions_book["Pending Month Conflict"]
    pending_headers = {cell.value: cell.column for cell in pending_sheet[2]}
    assert pending_sheet.cell(3, pending_headers["Reporting Month"]).value is None
    assert pending_sheet.cell(3, pending_headers["Reporting Month Status"]).value == "UNRESOLVED"
    exceptions_book.close()


def test_exported_null_credit_is_blank_and_zero_credit_is_numeric_zero(tmp_path) -> None:
    results = [
        _result(
            client="CLIENT-001", row=1, credit=None, currency=None,
            status="EXCEPTION", match_status="UNMATCHED", confidence="LOW",
            exception_codes=("NO_FINANCE_MATCH",),
        ),
        _result(
            client="CLIENT-001", row=2, credit=Decimal("0"), currency="EUR",
            status="ZERO_CREDIT",
        ),
    ]
    bundle = export_phase10_workbooks(
        build_detailed_report(results),
        build_executive_summary(results),
        build_client_totals(results),
        build_exceptions_report(results),
        tmp_path,
    )

    workbook = load_workbook(bundle.paths["CLIENT-001_APR26_Credit_Report.xlsx"])
    worksheet = workbook["Credit Detail"]
    headers = {cell.value: cell.column for cell in worksheet[2]}
    amount_column = headers["Calculated Credit"]
    assert worksheet.cell(3, amount_column).value is None
    assert worksheet.cell(4, amount_column).value == 0
    workbook.close()


def test_client_summary_sheets_are_inserted_first_without_rewriting_management_reports(
    tmp_path,
) -> None:
    unknown_total = Decimal(
        "21469.4194183928702104404702486177512721127279580188"
    )
    results = [
        _result(
            client="CLIENT-001", row=1, credit=Decimal("1084.716168925"), currency="EUR"
        ),
        _result(
            client="CLIENT-002", row=2, credit=Decimal("1240.742913354"), currency="USD"
        ),
        _result(
            client="CLIENT-004", row=3, credit=unknown_total, currency=None,
            rule="CLIENT_004_SPECIAL", match_status="UNMATCHED", confidence="LOW",
            exception_codes=("NO_FINANCE_MATCH",),
        ),
    ]
    detailed = build_detailed_report(results)
    client_totals = build_client_totals(results)
    executive = build_executive_summary(results)
    exceptions = build_exceptions_report(results)
    bundle = export_phase10_workbooks(
        detailed,
        executive,
        client_totals,
        exceptions,
        tmp_path,
    )
    executive_path = bundle.paths["Executive_Report.xlsx"]
    approved_executive_bytes = executive_path.read_bytes()

    added_paths = add_client_executive_summaries(detailed, client_totals, tmp_path)

    assert executive_path.read_bytes() == approved_executive_bytes
    for client in ("CLIENT-001", "CLIENT-002", "CLIENT-004"):
        path = added_paths[f"{client}_APR26_Credit_Report.xlsx"]
        workbook = load_workbook(path, data_only=False)
        assert workbook.sheetnames[0] == "Executive Summary"
        assert workbook.sheetnames[1:] == ["Credit Detail", "Notes & Methodology"]
        summary = workbook["Executive Summary"]
        summary_values = {
            summary.cell(row, 1).value: summary.cell(row, 2).value
            for row in range(4, summary.max_row + 1)
        }
        section_row = next(
            row
            for row in range(1, summary.max_row + 1)
            if summary.cell(row, 1).value == "BREACHED MARKETS ONLY (SLA BREACH % > 0)"
        )
        header_row = section_row + 1
        headers = {
            summary.cell(header_row, column).value: column
            for column in range(1, summary.max_column + 1)
            if summary.cell(header_row, column).value is not None
        }
        breached_rows = [
            row
            for row in range(header_row + 1, summary.max_row + 1)
            if summary.cell(row, 1).value is not None
        ]
        assert len(breached_rows) == 1
        confidence = summary.cell(breached_rows[0], headers["Match Confidence"]).value
        assert confidence in {"HIGH", "MEDIUM", "LOW"}
        assert summary_values["Total Markets Reviewed"] == 1
        assert summary_values["Markets with SLA Breach"] == 1
        assert sum(
            1 for _ in workbook["Credit Detail"].iter_rows(min_row=3, values_only=True)
        ) == 1

        if client == "CLIENT-004":
            assert summary_values["Currency / Currency Status"] == "UNKNOWN / UNSPECIFIED"
            assert summary_values["Total Confirmed Calculated Credit"] == (
                "21,469.42 - Currency Unspecified"
            )
            assert confidence == "LOW"
            assert summary.cell(breached_rows[0], headers["Match Status"]).value == "UNMATCHED"
            assert summary.cell(breached_rows[0], headers["Calculation Status"]).value == "CALCULATED"
            assert summary.cell(
                breached_rows[0], headers["Calculated Credit Receivable"]
            ).value == str(unknown_total)
        workbook.close()