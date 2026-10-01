"""Offline integration controls over the supplied assessment workbooks."""

from __future__ import annotations

import ast
import os
import socket
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.audit import run_phase12_reconciliations
from src.excel_export import (
	add_client_executive_summaries,
	verify_client_executive_summaries,
)
from src.ingestion import load_all_devops_files, load_finance_file
from src.matching import match_sla_record
from src.models import SLARecord
from src.normalization import normalize_devops_records, normalize_finance_records
from src.processing import apply_composite_allocation_control, process_sla_records
from src.reporting import (
	UNKNOWN_CURRENCY_BUCKET,
	build_client_totals,
	build_detailed_report,
	build_exceptions_report,
	build_executive_summary,
	reconcile_report_datasets,
)
from src.validation import validate_devops, validate_finance


ROOT = Path(__file__).resolve().parents[1]


def _run_assessment_inputs():
	finance_raw = load_finance_file(
		ROOT / "input/finance/00_Anonymized_Finance_APR26.xlsx"
	)
	devops_raw = load_all_devops_files(ROOT / "input/devops")
	finance = normalize_finance_records(finance_raw)
	devops = normalize_devops_records(devops_raw)
	matches = [match_sla_record(row, finance) for _, row in devops.iterrows()]
	slas = [
		SLARecord(
			client=row.get("Client"),
			raw_market_code=row.get("Market Code"),
			market_name=row.get("Market Name"),
			market_description=row.get("Market Description"),
			sla_breach_pct=row.get("SLA Breach %"),
			source_file=row.get("source_file"),
			source_sheet=row.get("source_sheet"),
			source_row_number=int(row.get("source_row_number")),
			filename_month_hint=row.get("filename_month_hint"),
			workbook_month_hint=row.get("workbook_month_hint"),
			month_conflict=bool(row.get("month_conflict")),
		)
		for _, row in devops.iterrows()
	]
	source_exceptions = tuple(
		validate_finance(finance_raw) + validate_devops(devops_raw)
	)
	normalized_markets = devops["normalized_market_code"].tolist()
	matches = apply_composite_allocation_control(slas, matches, normalized_markets)
	results = process_sla_records(
		slas,
		matches,
		normalized_devops_markets=normalized_markets,
		source_exceptions=source_exceptions,
	)
	detailed = build_detailed_report(results)
	client_totals = build_client_totals(results)
	executive = build_executive_summary(results)
	exceptions = build_exceptions_report(results)
	controls = reconcile_report_datasets(
		results, detailed, client_totals, executive, exceptions
	)
	return devops_raw, slas, matches, results, detailed, client_totals, executive, exceptions, controls


def _business_result_fingerprint(results):
	return tuple(
		(
			result.client,
			result.devops_market,
			result.normalized_devops_market,
			result.match_status,
			result.match_confidence,
			result.match_method,
			result.candidate_count,
			result.sla_breach_pct,
			result.applied_rate,
			result.base_credit,
			result.currency,
			result.calculation_rule,
			result.calculated_credit,
			result.status,
			result.exception_code,
			result.exception_codes,
			result.reporting_month,
			result.reporting_month_status,
		)
		for result in results
	)


def _executive_amount(frame, client, currency):
	rows = frame.loc[
		frame["Client"].eq(client) & frame["Currency"].eq(currency),
		"Total Calculated Credit",
	]
	assert len(rows) == 1
	return rows.iloc[0]


def test_actual_assessment_regression_offline_reproducible_and_phase10_complete(
	monkeypatch,
	tmp_path,
) -> None:
	for name in tuple(os.environ):
		if name.startswith("AWS_") or "BEDROCK" in name.upper():
			monkeypatch.delenv(name, raising=False)

	def deny_network(*_args, **_kwargs):
		raise AssertionError("Deterministic regression tests must not access the network.")

	monkeypatch.setattr(socket.socket, "connect", deny_network)
	monkeypatch.setattr(socket, "create_connection", deny_network)

	first = _run_assessment_inputs()
	second = _run_assessment_inputs()
	devops_raw, slas, matches, results, detailed, client_totals, executive, exceptions, controls = first
	_, _, second_matches, second_results, second_detail, second_totals, second_executive, _, second_controls = second

	assert len(devops_raw) == len(results) == len(detailed) == 244
	assert detailed["Reporting Month Status"].eq("RESOLVED").sum() == 178
	assert detailed["Reporting Month Status"].eq("UNRESOLVED").sum() == 66
	assert controls.passed and second_controls.passed
	assert _business_result_fingerprint(results) == _business_result_fingerprint(second_results)
	assert [
		(match.status, match.confidence, match.match_method, match.candidate_count, match.exception_code)
		for match in matches
	] == [
		(match.status, match.confidence, match.match_method, match.candidate_count, match.exception_code)
		for match in second_matches
	]
	pd.testing.assert_frame_equal(
		detailed.drop(columns=["Processing Run ID"]),
		second_detail.drop(columns=["Processing Run ID"]),
	)
	pd.testing.assert_frame_equal(client_totals, second_totals)
	pd.testing.assert_frame_equal(executive, second_executive)

	expected = {
		"CLIENT-001": (56, 7, 2, 5, 0),
		"CLIENT-002": (66, 8, 2, 4, 2),
		"CLIENT-004": (56, 4, 4, 0, 4),
	}
	for client, values in expected.items():
		rows = [
			result for result in results
			if result.client == client and result.reporting_month == "2026-04"
		]
		breached = [
			result for result in rows
			if Decimal(str(result.sla_breach_pct)) > 0
		]
		calculated_breached = [
			result for result in breached
			if result.status in {"CALCULATED", "ZERO_CREDIT"}
			and result.calculated_credit is not None
		]
		ambiguous = [result for result in breached if result.match_status == "AMBIGUOUS"]
		unmatched = [result for result in breached if result.match_status == "UNMATCHED"]
		assert (len(rows), len(breached), len(calculated_breached), len(ambiguous), len(unmatched)) == values

	allocation_rows = {
		("CLIENT-001", 2): ("NYMEX", 36),
		("CLIENT-001", 3): ("CBOT", 36),
		("CLIENT-001", 5): ("CME", 36),
		("CLIENT-002", 4): ("NYMEX", 1112),
		("CLIENT-002", 5): ("CBOT", 1112),
		("CLIENT-002", 7): ("client002_lifams", 1152),
		("CLIENT-002", 9): ("client002_lifpar", 1152),
		("CLIENT-002", 49): ("client002_lifbru", 1152),
	}
	allocation_results = {
		(result.client, result.source_row_number): result
		for result in results
		if result.exception_code == "COMPOSITE_ALLOCATION_AMBIGUITY"
	}
	assert set(allocation_results) == set(allocation_rows)
	for key, (market, finance_row) in allocation_rows.items():
		result = allocation_results[key]
		assert result.normalized_devops_market == market.upper().replace("CLIENT002_", "") or result.devops_market == market
		assert result.match_status == "AMBIGUOUS"
		assert result.match_confidence == "LOW"
		assert result.status == "EXCEPTION"
		assert result.calculated_credit is None
		assert result.applied_rate is None
		assert result.base_credit is None
		assert result.currency is None
		assert result.finance_source_row_number is None
		assert len(result.candidate_finance_records) == 1
		assert result.candidate_finance_records[0].source_row_number == finance_row
		assert "No allocation rule is provided" in result.explanation
	assert _executive_amount(executive, "CLIENT-001", "EUR") == Decimal("702.960636375")
	assert _executive_amount(executive, "CLIENT-002", "USD") == Decimal("463.360430778")

	client001_tse = next(
		result for result in results
		if result.client == "CLIENT-001" and result.normalized_devops_market == "TSE"
	)
	assert (client001_tse.match_status, client001_tse.match_confidence, client001_tse.candidate_count) == (
		"AMBIGUOUS",
		"LOW",
		2,
	)
	for market in ("CME", "ICEUS", "TRQEE"):
		result = next(
			item for item in results
			if item.client == "CLIENT-002" and item.normalized_devops_market == market
		)
		assert result.match_status == "AMBIGUOUS"
		assert result.match_confidence == "LOW"
		assert result.calculated_credit is None
	client002_iceeu = next(
		result for result in results
		if result.client == "CLIENT-002" and result.normalized_devops_market == "ICEEU"
	)
	assert client002_iceeu.match_status == "AMBIGUOUS"
	assert client002_iceeu.candidate_count == 3
	assert _executive_amount(
		executive,
		"CLIENT-004",
		UNKNOWN_CURRENCY_BUCKET,
	) == Decimal("21469.4194183928702104404702486177512721127279580188")
	assert set(executive["Reporting Month"]) == {"2026-04"}
	assert sum(result.exception_code == "MONTH_CONFLICT" for result in results) == 66

	client004_breached = [
		result for result in results
		if result.client == "CLIENT-004"
		and result.reporting_month == "2026-04"
		and Decimal(str(result.sla_breach_pct)) > 0
	]
	assert len(client004_breached) == 4
	assert all(
		result.match_status == "UNMATCHED"
		and result.match_confidence == "LOW"
		and result.status == "CALCULATED"
		and result.currency is None
		and result.base_credit is None
		for result in client004_breached
	)
	assert all(
		result.calculated_credit is None
		for result in results
		if result.calculation_rule == "STANDARD_CAPPED"
		and result.match_status in {"AMBIGUOUS", "UNMATCHED"}
	)
	assert all(
		result.source_file is None or "MAY26" not in result.source_file
		for result in results
		if result.reporting_month == "2026-04"
	)

	from src.excel_export import add_client_executive_summaries, export_phase10_workbooks, verify_client_executive_summaries

	export = export_phase10_workbooks(
		detailed,
		executive,
		client_totals,
		exceptions,
		tmp_path,
	)
	paths = add_client_executive_summaries(detailed, client_totals, tmp_path)
	workbook_checks = verify_client_executive_summaries(paths, detailed)
	assert all(all(checks.values()) for checks in workbook_checks.values())
	assert len(export.paths) == 5
	phase12 = run_phase12_reconciliations(
		slas,
		matches,
		results,
		detailed,
		client_totals,
		executive,
		exceptions,
		output_dir=tmp_path,
	)
	assert phase12.passed, "\n".join(
		f"{control.control}: expected={control.expected}; actual={control.actual}; "
		f"status={control.status}; {control.explanation}"
		for control in phase12.controls
		if control.status == "FAIL"
	)
	for client, expected_count in (("CLIENT-001", 7), ("CLIENT-002", 8), ("CLIENT-004", 4)):
		workbook = load_workbook(paths[f"{client}_APR26_Credit_Report.xlsx"], data_only=False)
		assert workbook.sheetnames == ["Executive Summary", "Credit Detail", "Notes & Methodology"]
		summary = workbook["Executive Summary"]
		assert summary.auto_filter.ref
		metrics = {
			summary.cell(row, 1).value: summary.cell(row, 2).value
			for row in range(4, summary.max_row + 1)
		}
		expected_summary = {
			"CLIENT-001": (
				"EUR",
				"EUR 702.96",
				Decimal("702.960636375"),
			),
			"CLIENT-002": (
				"USD",
				"USD 463.36",
				Decimal("463.360430778"),
			),
			"CLIENT-004": (
				"UNKNOWN / UNSPECIFIED",
				"21,469.42 - Currency Unspecified",
				Decimal("21469.4194183928702104404702486177512721127279580188"),
			),
		}[client]
		assert metrics["Currency / Currency Status"] == expected_summary[0]
		assert metrics["Total Confirmed Calculated Credit"] == expected_summary[1]
		if client == "CLIENT-004":
			assert metrics["Authoritative Total (Exact)"] == format(expected_summary[2], "f")
		else:
			assert Decimal(str(metrics["Authoritative Total (Exact)"])) == expected_summary[2]
		assert workbook["Credit Detail"].max_row - 2 == expected[client][0]
		workbook.close()


def test_deterministic_core_has_no_aws_imports_and_lambda_sdk_is_lazy() -> None:
	forbidden_roots = {"aws", "awscrt", "bedrock", "boto3", "botocore", "langchain"}
	imports = []
	for source_path in sorted((ROOT / "src").glob("*.py")):
		if source_path.name == "lambda_handler.py":
			continue
		tree = ast.parse(source_path.read_text(encoding="utf-8"))
		for node in ast.walk(tree):
			if isinstance(node, ast.Import):
				imports.extend(alias.name.split(".")[0] for alias in node.names)
			elif isinstance(node, ast.ImportFrom) and node.module:
				imports.append(node.module.split(".")[0])
	assert not forbidden_roots.intersection(imports)

	lambda_tree = ast.parse(
		(ROOT / "src/lambda_handler.py").read_text(encoding="utf-8")
	)
	loader = next(
		node for node in lambda_tree.body
		if isinstance(node, ast.FunctionDef) and node.name == "_load_pinned_s3_snapshot"
	)
	assert any(
		isinstance(node, ast.Import)
		and any(alias.name.split(".")[0] == "boto3" for alias in node.names)
		for node in ast.walk(loader)
	)
	assert not any(
		isinstance(node, ast.Import)
		and any(alias.name.split(".")[0] == "boto3" for alias in node.names)
		for node in lambda_tree.body
	)